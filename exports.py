"""Download-only supplier files. This module cannot transmit orders."""
import csv
import io
from decimal import Decimal, ROUND_HALF_UP
from localization import message, normalize_locale


COLUMNS = [
    ("status", "Статус"), ("supplier", "Поставщик"), ("warehouse", "Склад"),
    ("sku", "Артикул"), ("name", "Наименование"), ("unit", "Ед. изм."),
    ("recommended_quantity", "Рекомендовано"), ("quantity", "Количество к заказу"),
    ("unit_price", "Цена (из источника)"), ("amount", "Сумма (из источника)"),
    ("urgency", "Срочность"), ("explanation", "Обоснование"),
    ("reviewer", "Утвердил"), ("approved_at", "Дата утверждения"),
]


EXPORT_TEXT = {
    "ru": {
        "headers": [title for _, title in COLUMNS],
        "approved": "Утверждён", "draft": "Черновик — требует проверки",
        "orders_sheet": "Заказы поставщикам", "methodology_sheet": "Методология",
        "field": "Поле", "value": "Значение", "run": "Расчёт", "as_of": "На дату",
        "source": "Источник", "status": "Статус", "purpose": "Назначение", "limitation": "Ограничение",
        "purpose_text": "Файл для проверки и ручной загрузки. Автоматическая отправка поставщику не выполняется.",
        "unknown_price": "Пустые цена и сумма означают отсутствие подтверждённой цены в источнике.",
    },
    "kk": {
        "headers": ["Мәртебе", "Жеткізуші", "Қойма", "Артикул", "Атауы", "Өлшем бірлігі",
                    "Ұсынылған саны", "Тапсырыс саны", "Баға (дереккөзден)", "Сома (дереккөзден)",
                    "Шұғылдық", "Негіздеме", "Бекіткен тұлға", "Бекітілген күні"],
        "approved": "Бекітілген", "draft": "Жоба — тексеруді қажет етеді",
        "orders_sheet": "Жеткізушілерге тапсырыстар", "methodology_sheet": "Әдістеме",
        "field": "Өріс", "value": "Мәні", "run": "Есептеу", "as_of": "Есептеу күні",
        "source": "Дереккөз", "status": "Мәртебе", "purpose": "Мақсаты", "limitation": "Шектеу",
        "purpose_text": "Тексеруге және қолмен жүктеуге арналған файл. Жеткізушіге автоматты түрде жіберілмейді.",
        "unknown_price": "Баға мен сома бос болса, дереккөзде расталған баға жоқ.",
    },
    "en": {
        "headers": ["Status", "Supplier", "Warehouse", "SKU", "Name", "Unit",
                    "Recommended quantity", "Order quantity", "Price (from source)", "Amount (from source)",
                    "Urgency", "Rationale", "Approved by", "Approval date"],
        "approved": "Approved", "draft": "Draft — review required",
        "orders_sheet": "Supplier orders", "methodology_sheet": "Methodology",
        "field": "Field", "value": "Value", "run": "Calculation", "as_of": "As of",
        "source": "Source", "status": "Status", "purpose": "Purpose", "limitation": "Limitation",
        "purpose_text": "For review and manual import. Orders are not sent to suppliers automatically.",
        "unknown_price": "Blank price and amount mean that the source does not provide a confirmed price.",
    },
}


def export_locale(run):
    """Use the saved calculation language; older runs retain Russian defaults."""
    return normalize_locale(run.get("locale", (run.get("settings") or {}).get("locale", "ru")))


def row_key(row):
    return f"{row.get('warehouse', '')}::{row['sku']}"


def export_rows(run, supplier=None):
    locale = export_locale(run)
    text = EXPORT_TEXT[locale]
    approval = run.get("approval")
    rows = []
    for source in run["rows"]:
        if supplier and source.get("supplier") != supplier:
            continue
        row = dict(source)
        quantity = approval["quantities"][row_key(row)] if approval else row["recommended_quantity"]
        row.update(quantity=quantity, status=text["approved"] if approval else text["draft"],
                   reviewer=approval["reviewer"] if approval else "",
                   approved_at=approval["approved_at"] if approval else "")
        if row.get("urgency") in {"critical", "high", "normal", "covered"}:
            row["urgency"] = message(locale, f"urgency_{row['urgency']}")
        price_known = row.get("price_known", row.get("unit_price") not in (None, 0, ""))
        if price_known:
            row["amount"] = float((Decimal(str(quantity)) * Decimal(str(row["unit_price"]))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
        else:
            # Unknown is not a confirmed zero price. Leave the business cells empty.
            row["unit_price"] = ""
            row["amount"] = ""
        rows.append(row)
    return rows


def safe_cell(value):
    # Prevent untrusted names/identifiers becoming formulas in spreadsheet viewers.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value if value is not None else ""


def to_csv(run, supplier=None):
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(EXPORT_TEXT[export_locale(run)]["headers"])
    for row in export_rows(run, supplier):
        writer.writerow([safe_cell(row.get(key, "")) for key, _ in COLUMNS])
    return output.getvalue().encode("utf-8-sig")


def to_xlsx(run, supplier=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    text = EXPORT_TEXT[export_locale(run)]
    book = Workbook()
    sheet = book.active
    sheet.title = text["orders_sheet"]
    sheet.append(text["headers"])
    for row in export_rows(run, supplier):
        sheet.append([safe_cell(row.get(key, "")) for key, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="15263A")
        cell.alignment = Alignment(wrap_text=True)
    widths = [28, 25, 22, 24, 44, 12, 18, 24, 20, 20, 16, 80, 24, 28]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (5, 12))
        for index in (7, 8, 9, 10):
            row[index-1].number_format = "#,##0.00"
    info = book.create_sheet(text["methodology_sheet"])
    info.append([text["field"], text["value"]])
    info.append([text["run"], safe_cell(run["run_id"])])
    info.append([text["as_of"], safe_cell(run.get("as_of", ""))])
    info.append([text["source"], safe_cell(run.get("source_label", ""))])
    info.append([text["status"], text["approved"] if run.get("approval") else text["draft"]])
    info.append([text["purpose"], text["purpose_text"]])
    if any(row.get("price_known") is False for row in run["rows"]):
        info.append([text["limitation"], text["unknown_price"]])
    for warning in run.get("warnings", []):
        info.append([text["limitation"], safe_cell(str(warning))])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 100
    for row in info:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
