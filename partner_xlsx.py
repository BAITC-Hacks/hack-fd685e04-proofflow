"""Local-only adapter for the partner's IEK and Systeme Electric workbooks.

The adapter does not infer a customer identifier from a document number, does
not invent prices or lead times, and never transmits workbook contents.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import re
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from openpyxl import load_workbook


def _day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return None
    text = str(value).strip()
    if re.match(r"^\d{2}\.\d{2}\.\d{4}", text):
        try:
            return datetime.strptime(text[:10], "%d.%m.%Y").date().isoformat()
        except ValueError:
            return None
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return None


def _number(value):
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace("\u00a0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def _supplier(path: Path):
    name = str(path).casefold()
    if "iek" in name or "иэк" in name:
        return "IEK"
    if "systeme" in name or "syseme" in name or "сэ" in name:
        return "Systeme Electric"
    return None


def import_partner_files(paths: list[str]) -> dict:
    """Map local partner XLSX/ZIP files to the canonical procurement dataset."""
    with TemporaryDirectory(prefix="prooflow_partner_") as temp:
        files: list[Path] = []
        aggregate_bytes = 0
        for raw_path in paths:
            source = Path(raw_path)
            if source.suffix.lower() == ".xlsx":
                files.append(source)
            elif source.suffix.lower() == ".zip":
                with ZipFile(source) as archive:
                    if len(archive.infolist()) > 100:
                        raise ValueError("Partner ZIP has too many members")
                    for member in archive.infolist():
                        if member.is_dir() or not member.filename.lower().endswith(".xlsx"):
                            continue
                        if Path(member.filename).name.startswith("~$"):
                            continue
                        if member.file_size > 80 * 1024 * 1024:
                            raise ValueError("Partner workbook exceeds 80 MB")
                        aggregate_bytes += member.file_size
                        if aggregate_bytes > 160 * 1024 * 1024 or member.file_size > max(member.compress_size, 1) * 100:
                            raise ValueError("Partner ZIP expanded size exceeds safety limit")
                        # Use the basename only: archive member paths must not escape temp.
                        supplier_hint = _supplier(source) or _supplier(Path(member.filename)) or "unknown"
                        target = Path(temp) / f"{len(files)}_{supplier_hint}_{Path(member.filename).name}"
                        target.write_bytes(archive.read(member))
                        files.append(target)
                        if len(files) > 24:
                            raise ValueError("Select at most 24 partner workbooks")
        if len(files) > 24:
            raise ValueError("Select at most 24 partner workbooks")
        return _read_workbooks(files)


def _read_workbooks(files: list[Path]) -> dict:
    products = {}
    sales = []
    inbound = []
    warnings = set()
    counts = defaultdict(int)
    total_xlsx_expanded = 0
    for path in files:
        supplier = _supplier(path)
        if not supplier:
            warnings.add(f"Не распознан поставщик файла {path.name}; файл пропущен.")
            continue
        # XLSX is itself a ZIP. Bound decompression before openpyxl parses XML.
        with ZipFile(path) as internal:
            members = internal.infolist()
            expanded = sum(member.file_size for member in members)
            total_xlsx_expanded += expanded
            if len(members) > 128 or total_xlsx_expanded > 512 * 1024 * 1024:
                raise ValueError("Partner XLSX collection exceeds safety limit")
            if expanded > 160 * 1024 * 1024 or any(
                member.file_size > max(member.compress_size, 1) * 100
                for member in members
            ):
                raise ValueError("XLSX expanded size exceeds safety limit")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            _read_one(path, workbook, supplier, products, sales, inbound, warnings, counts)
        finally:
            workbook.close()
    if not sales:
        raise ValueError("Не найдены дневные продажи в предоставленных книгах")
    warnings.update({
        "В исходных продажах нет ID клиента: номер документа — только event_id.",
        "Цена продажи и подтверждённый срок поставки отсутствуют; заполните параметры перед утверждением.",
        "Точных периодов stockout нет; компенсация потерянного спроса доступна после загрузки этих интервалов.",
    })
    if counts["nonpositive_sales"]:
        warnings.add(f"Исключено {counts['nonpositive_sales']} неположительных строк продаж: возвраты/корректировки требуют сверки.")
    if counts["invalid_sales"]:
        warnings.add(f"Пропущено {counts['invalid_sales']} неполных или некорректных строк продаж.")
    return {"as_of": max(row["date"] for row in sales),
            "products": list(products.values()), "sales": sales,
            "inbound": inbound, "stockouts": [], "settings": {},
            "category_policies": {},
            "metadata": {"source_label": "Локальные файлы Электрокомплекта",
                         "synthetic": False, "warnings": sorted(warnings),
                         "record_counts": dict(counts)}}


def _read_one(path, workbook, supplier, products, sales, inbound, warnings, counts):
    """Dispatch by documented source-sheet identity; unknown books stay inert."""
    name = path.name.casefold()
    if "динамика продаж" in name:
        _read_sales(workbook.worksheets[0], supplier, products, sales, counts)
    elif "остатки" in name:
        _read_monthly_stock(workbook.worksheets[0], supplier, products, warnings, counts)
    elif "moq" in name:
        _read_pack(workbook.worksheets[0], supplier, products, counts)
    elif "путь" in name and supplier == "IEK":
        _read_iek_inbound(workbook.worksheets[0], products, inbound, warnings, counts)
    elif "пути" in name and supplier == "Systeme Electric":
        _read_systeme_report(workbook.worksheets[0], products, inbound, warnings, counts)
    elif "сезонность" in name or "ежемесячные продажи" in name:
        counts["crosscheck_workbooks"] += 1
        warnings.add("Месячные продажи/сезонность не суммируются с дневными продажами во избежание двойного счёта.")
    else:
        counts["unmapped_workbooks"] += 1
        warnings.add(f"Лист {path.name} пока не сопоставлен: его значения не включены в расчёт.")


def _read_sales(sheet, supplier, products, sales, counts):
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)
    for values in rows:
        if len(values) < 8:
            counts["invalid_sales"] += 1
            continue
        day, document_number, _, sku, name, unit, warehouse, quantity = values[:8]
        day = _day(day)
        sku = str(sku).strip() if sku is not None else ""
        qty = _number(quantity)
        if not day or not sku or qty is None:
            counts["invalid_sales"] += 1
            continue
        if qty <= 0:
            counts["nonpositive_sales"] += 1
            continue
        warehouse = str(warehouse or "Алматы").strip()
        key = warehouse, sku
        if key not in products:
            _update_product(products, sku, supplier, name=name, unit=unit,
                            warehouse=warehouse)
        else:
            # The sales sheet carries the measured unit. It can follow a pack
            # or inventory sheet that only provided a generic placeholder.
            products[key]["unit"] = str(unit or products[key]["unit"])
        event = {"date": day, "sku": sku, "warehouse": warehouse, "quantity": qty}
        if document_number not in (None, ""):
            event["event_id"] = str(document_number)
        sales.append(event)
        counts["sales"] += 1


def _update_product(products, sku, supplier, *, name=None, unit=None, warehouse="Алматы"):
    # The 1C code, not the supplier article, is the common key across sheets.
    sku = str(sku or "").strip()
    if not sku:
        return None
    key = warehouse, sku
    if key not in products:
        products[key] = {"sku": sku, "warehouse": warehouse, "name": str(name or sku),
                         "supplier": supplier, "category": "unknown", "unit": str(unit or "pcs"),
                         "on_hand": 0, "lead_days": 0, "moq": 0, "pack_size": 1,
                         "unit_price": 0,
                         "provenance": {"on_hand": "missing", "lead_days": "missing",
                                        "unit_price": "missing"}}
    return products[key]


def _read_monthly_stock(sheet, supplier, products, warnings, counts):
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, ())
    sku_col = 2 if supplier == "IEK" else 2
    name_col = 0 if supplier == "IEK" else 1
    unit_col = 1 if supplier == "IEK" else 3
    # Last dated September 2026 column, never the undated “Итого” summary.
    last_col = next((index for index, value in enumerate(headers)
                     if str(value or "").casefold().startswith("сент") and "2026" in str(value)), None)
    if last_col is None:
        warnings.add(f"Не найден датированный остаток в {supplier}; остаток не импортирован.")
        return
    for values in rows:
        if len(values) <= max(sku_col, last_col):
            continue
        qty = _number(values[last_col])
        if qty is None:
            continue
        product = _update_product(products, values[sku_col], supplier,
                                  name=values[name_col], unit=values[unit_col])
        if product is None:
            continue
        if qty < 0:
            counts["negative_stock"] += 1
            continue
        # Systeme's dedicated current stock report has greater specificity.
        if product["provenance"]["on_hand"] == "observed_2026-09-22_report":
            continue
        product["on_hand"] = qty
        product["provenance"]["on_hand"] = "monthly_2026-09_provisional"
        counts["monthly_stock"] += 1
    warnings.add("Остатки месячной матрицы — датированный сентябрьский срез, не подтверждённый складской остаток на 22.09; менеджер должен сверить его перед утверждением.")


def _read_pack(sheet, supplier, products, counts):
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)
    sku_col = 1 if supplier == "IEK" else 2
    pack_col = 4
    name_col = 3 if supplier == "IEK" else 1
    for values in rows:
        if len(values) <= pack_col:
            continue
        pack = _number(values[pack_col])
        if pack is None or pack <= 0:
            continue
        product = _update_product(products, values[sku_col], supplier, name=values[name_col])
        if product is not None:
            product["pack_size"] = pack
            product["provenance"]["pack_size"] = "partner_shipping_multiple"
            counts["pack_size"] += 1


def _read_iek_inbound(sheet, products, inbound, warnings, counts):
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, ())
    eta_by_col = {}
    for index, title in enumerate(headers):
        if index < 3:
            continue
        match = re.search(r"(?:до\s+)(\d{2})\.(\d{2})\.(\d{4})", str(title or ""), re.IGNORECASE)
        if match:
            d, m, y = map(int, match.groups())
            try:
                eta_by_col[index] = date(y, m, d).isoformat()
            except ValueError:
                continue
    if not eta_by_col:
        warnings.add("IEK: не распознаны ETA партий; поступления не импортированы.")
    for values in rows:
        if not values:
            continue
        product = _update_product(products, values[0], "IEK", name=values[2] if len(values) > 2 else None)
        if product is None:
            continue
        for index, eta in eta_by_col.items():
            qty = _number(values[index]) if len(values) > index else None
            if qty is not None and qty > 0:
                inbound.append({"sku": product["sku"], "warehouse": "Алматы",
                                "quantity": qty, "eta": eta})
                counts["inbound"] += 1
    warnings.add("IEK: ETA в заголовках означает плановое поступление «до» даты; фактическая дата/поставка не подтверждены.")


def _read_systeme_report(sheet, products, inbound, warnings, counts):
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)  # decorative first row
    headers = next(rows, ())
    labels = {str(value or "").strip().casefold(): index for index, value in enumerate(headers)}
    def col(name):
        return labels.get(name.casefold())
    sku_col = col("Код 1с")
    free_col = col("Свободный остаток")
    category_col = col("Категория 2026")
    inbound_col = next((index for label, index in labels.items() if "в пути 24.09" in label), None)
    if sku_col is None or free_col is None:
        warnings.add("Systeme: схема отчёта TDSheet изменилась; остатки не импортированы.")
        return
    for values in rows:
        if len(values) <= sku_col:
            continue
        product = _update_product(products, values[sku_col], "Systeme Electric",
                                  name=values[col("Наименование")] if col("Наименование") is not None else None)
        if product is None:
            continue
        if category_col is not None and values[category_col] not in (None, ""):
            product["category"] = str(values[category_col]).strip()
            product["provenance"]["category"] = "partner_2026_category"
            counts["category"] += 1
        stock = _number(values[free_col]) if len(values) > free_col else None
        if stock is not None and stock >= 0:
            product["on_hand"] = stock
            product["provenance"]["on_hand"] = "observed_2026-09-22_report"
            counts["report_stock"] += 1
        elif stock is not None:
            counts["negative_stock"] += 1
        if inbound_col is not None and len(values) > inbound_col:
            qty = _number(values[inbound_col])
            if qty is not None and qty > 0:
                inbound.append({"sku": product["sku"], "warehouse": "Алматы",
                                "quantity": qty, "eta": "2026-09-24"})
                counts["inbound"] += 1
    warnings.add("Systeme: «Свободный остаток» принят за доступный агрегированный остаток; складское распределение и смысл «в пути 24.09» требуют подтверждения.")
