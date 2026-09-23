"""Repeatable full synthetic procurement journey through the real API/engine."""
import io
import json
import os
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from server import app


class EndToEndTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {"PROOFFLOW_DATA_DIR": self.temp.name})
        self.env.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.env.stop()
        self.temp.cleanup()

    def test_demo_calculate_explain_adjust_approve_export(self):
        demo = self.client.post("/api/demo")
        self.assertEqual(demo.status_code, 200, demo.text[:500])
        dataset = demo.json()["dataset"]
        self.assertTrue(dataset["metadata"]["synthetic"])
        response = self.client.post("/api/calculate", json={"settings": {"review_days": 14}})
        self.assertEqual(response.status_code, 200, response.text[:500])
        calculation = response.json()
        self.assertGreaterEqual(len(calculation["rows"]), 5)
        for row in calculation["rows"]:
            self.assertTrue(row["explanation"])
            self.assertGreaterEqual(row["recommended_quantity"], 0)
        row = calculation["rows"][0]
        self.assertNotIn("history", row)
        detail = self.client.get(f"/api/runs/{calculation['run_id']}/item", params={"sku": row["sku"], "warehouse": row["warehouse"]})
        self.assertEqual(detail.status_code, 200)
        self.assertTrue(detail.json()["history"])
        row_key = row["warehouse"] + "::" + row["sku"]
        chosen = row["recommended_quantity"] + 5
        run_id = calculation["run_id"]
        approved = self.client.post(f"/api/runs/{run_id}/approve", json={"reviewer": "Synthetic test manager", "quantities": {row_key: chosen}})
        self.assertEqual(approved.status_code, 200, approved.text)
        self.assertEqual(approved.json()["quantities"][row_key], chosen)
        exported = self.client.get(f"/api/runs/{run_id}/export?format=xlsx")
        workbook = load_workbook(io.BytesIO(exported.content), data_only=True)
        sheet = workbook["Заказы поставщикам"]
        self.assertEqual(sheet.max_row - 1, len(calculation["rows"]))
        self.assertEqual(sheet["H2"].value, chosen)
        self.assertEqual(sheet["A2"].value, "Утверждён")
        workbook.close()

    def test_roundtrip_template_upload_and_supplier_filter(self):
        template = self.client.get("/api/template")
        self.assertEqual(template.status_code, 200)
        uploaded = self.client.post("/api/import", files={"files": ("synthetic.json", template.content, "application/json")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text[:500])
        supplier = uploaded.json()["dataset"]["products"][0]["supplier"]
        result = self.client.post("/api/calculate", json={"filters": {"supplier": supplier}})
        self.assertEqual(result.status_code, 200, result.text[:500])
        rows = result.json()["rows"]
        self.assertGreater(len(rows), 0)
        self.assertEqual({row["supplier"] for row in rows}, {supplier})

    def test_failed_import_does_not_replace_existing_dataset(self):
        self.client.post("/api/demo")
        before = self.client.get("/api/dataset").json()
        response = self.client.post("/api/import", files={"files": ("bad.json", b'{"broken":', "application/json")})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.get("/api/dataset").json(), before)

    def test_stale_preview_cannot_overwrite_new_dataset(self):
        old = self.client.post("/api/demo").json()["dataset"]
        self.assertTrue(old["_preview"])
        self.assertNotIn("sales", old)
        self.assertGreater(old["metadata"]["record_counts"]["sales"], 0)
        self.client.post("/api/demo")
        result = self.client.post("/api/calculate", json={"dataset": old})
        self.assertEqual(result.status_code, 409)

    def test_current_preview_product_edits_affect_calculation(self):
        preview = self.client.post("/api/demo").json()["dataset"]
        for product in preview["products"]:
            product["on_hand"] = 1_000_000_000
        response = self.client.post("/api/calculate", json={"dataset": preview})
        self.assertEqual(response.status_code, 200, response.text[:500])
        self.assertTrue(all(row["recommended_quantity"] == 0 for row in response.json()["rows"]))


if __name__ == "__main__":
    unittest.main()
