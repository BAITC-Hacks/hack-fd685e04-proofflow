"""Localized downloads preserve calculations, provenance and safe cells."""
from copy import deepcopy
import csv
import io

from openpyxl import load_workbook
import pytest

from exports import COLUMNS, EXPORT_TEXT, export_rows, to_csv, to_xlsx


def run_fixture():
    return {
        "run_id": "synthetic-export-i18n", "source_label": "=untrusted-source",
        "as_of": "2026-09-23", "warnings": ["@untrusted-warning"],
        "rows": [
            {"sku": "A", "name": "=untrusted-name", "warehouse": "W", "supplier": "S",
             "unit": "m", "recommended_quantity": 2.5, "unit_price": 1.25, "price_known": True,
             "urgency": "high", "explanation": "Source explanation"},
            {"sku": "B", "name": "B", "warehouse": "W", "supplier": "S", "unit": "pcs",
             "recommended_quantity": 3, "unit_price": 0, "price_known": False, "urgency": "normal"},
        ],
    }


@pytest.mark.parametrize("locale", ["ru", "kk", "en"])
@pytest.mark.parametrize("approved", [False, True])
def test_localized_downloads_preserve_numbers_and_open_in_excel(locale, approved):
    run = run_fixture()
    run["locale"] = locale
    if approved:
        run["approval"] = {"quantities": {"W::A": 4, "W::B": 1},
                           "reviewer": "Reviewer", "approved_at": "2026-09-23T12:00:00Z"}
    before = deepcopy(run)
    text = EXPORT_TEXT[locale]
    csv_rows = list(csv.reader(io.StringIO(to_csv(run).decode("utf-8-sig")), delimiter=";"))
    assert csv_rows[0] == text["headers"]
    assert csv_rows[1][0] == text["approved" if approved else "draft"]
    assert float(csv_rows[1][6]) == 2.5
    assert float(csv_rows[1][7]) == (4 if approved else 2.5)
    assert float(csv_rows[1][9]) == (5 if approved else 3.13)
    assert csv_rows[2][8:10] == ["", ""]
    assert csv_rows[1][4] == "'=untrusted-name"
    assert csv_rows[1][10] == {"ru": "высокая", "kk": "жоғары", "en": "high"}[locale]
    book = load_workbook(io.BytesIO(to_xlsx(run)), data_only=False)
    try:
        assert book.sheetnames == [text["orders_sheet"], text["methodology_sheet"]]
        sheet = book.worksheets[0]
        assert [cell.value for cell in sheet[1]] == text["headers"]
        assert sheet["G2"].value == 2.5
        assert sheet["H2"].value == (4 if approved else 2.5)
        assert sheet["J2"].value == (5 if approved else 3.13)
        assert sheet["I3"].value is None and sheet["J3"].value is None
        assert sheet["E2"].data_type == "s" and sheet["E2"].value.startswith("'")
        info = book.worksheets[1]
        assert info["B4"].value == "'=untrusted-source"
        assert info["B4"].data_type == "s"
        assert any(row[1].value == text["unknown_price"] for row in info)
        assert all(cell.data_type != "f" for row in info for cell in row)
    finally:
        book.close()
    assert run == before


def test_default_stays_russian_and_settings_locale_is_supported():
    run = run_fixture()
    rows = list(csv.reader(io.StringIO(to_csv(run).decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == [title for _, title in COLUMNS]
    run["settings"] = {"locale": "kk"}
    assert export_rows(run)[0]["status"] == "Жоба — тексеруді қажет етеді"


def test_numeric_order_rows_are_identical_across_languages():
    values = []
    for locale in ("ru", "kk", "en"):
        run = run_fixture()
        run["locale"] = locale
        values.append([{key: row[key] for key in ("sku", "recommended_quantity", "quantity", "unit_price", "amount")}
                       for row in export_rows(run)])
    assert values[0] == values[1] == values[2]
