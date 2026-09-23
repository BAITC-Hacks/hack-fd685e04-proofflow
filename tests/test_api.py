"""Local workflow contract tests; no network credentials or partner data."""
import csv
import io
import os
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import storage
from server import app


class APITests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {"PROOFFLOW_DATA_DIR": self.temp.name})
        self.env.start()
        self.client = TestClient(app)
        self.run = {
            "run_id": "synthetic-workflow-test", "as_of": "2026-09-23",
            "source_label": "Synthetic API fixture", "warnings": [],
            "rows": [{"sku": "TEST-001", "name": "Synthetic cable", "warehouse": "TEST-WH",
                      "supplier": "TEST-SUPPLIER", "category": "Cable", "unit": "m",
                      "recommended_quantity": 12, "unit_price": 2.5,
                      "urgency": "high", "explanation": "Demand 24 minus stock 12 = 12"}],
        }
        storage.save_run(self.run)

    def tearDown(self):
        self.client.close()
        self.env.stop()
        self.temp.cleanup()

    def test_health_and_empty_dataset(self):
        self.assertEqual(self.client.get("/api/health").json()["status"], "ok")
        self.assertIsNone(self.client.get("/api/dataset").json()["dataset"])

    def test_approval_preserves_recommendation_and_exports_manual_quantity(self):
        response = self.client.post("/api/runs/synthetic-workflow-test/approve", json={
            "reviewer": "Test manager", "quantities": {"TEST-WH::TEST-001": 16}})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertFalse(response.json()["supplier_sent"])
        stored = self.client.get("/api/runs/synthetic-workflow-test").json()
        self.assertEqual(stored["rows"][0]["recommended_quantity"], 12)
        self.assertEqual(stored["approval"]["quantities"]["TEST-WH::TEST-001"], 16)
        export = self.client.get("/api/runs/synthetic-workflow-test/export?format=csv")
        rows = list(csv.reader(io.StringIO(export.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(float(rows[1][7]), 16)
        self.assertEqual(float(rows[1][9]), 40)
        self.assertEqual(rows[1][0], "Утверждён")
        repeated = self.client.post("/api/runs/synthetic-workflow-test/approve", json={"reviewer": "Test manager"})
        self.assertEqual(repeated.status_code, 409)

    def test_draft_export_is_labelled(self):
        response = self.client.get("/api/runs/synthetic-workflow-test/export?format=xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        sheet = workbook["Заказы поставщикам"]
        self.assertIn("Черновик", sheet["A2"].value)
        self.assertEqual(sheet["H2"].value, 12)
        self.assertEqual(sheet["J2"].value, 30)
        workbook.close()

    def test_unknown_position_and_negative_quantity_are_rejected(self):
        for quantities in ({"missing": 1}, {"TEST-WH::TEST-001": -1}):
            response = self.client.post("/api/runs/synthetic-workflow-test/approve", json={
                "reviewer": "Test", "quantities": quantities})
            self.assertEqual(response.status_code, 422)
        self.assertIsNone(storage.load_run(self.run["run_id"])["approval"])

    def test_blank_reviewer_is_rejected(self):
        self.assertEqual(self.client.post("/api/runs/synthetic-workflow-test/approve", json={"reviewer": "  "}).status_code, 422)

    def test_missing_run_and_unsupported_export(self):
        self.assertEqual(self.client.get("/api/runs/nonexistent").status_code, 404)
        self.assertEqual(self.client.get("/api/runs/synthetic-workflow-test/export?format=exe").status_code, 422)

    def test_supplier_filter_matches_export(self):
        response = self.client.get("/api/runs/synthetic-workflow-test/export?supplier=unknown")
        self.assertEqual(len(list(csv.reader(io.StringIO(response.content.decode("utf-8-sig")), delimiter=";"))), 1)

    def test_external_origin_cannot_approve(self):
        response = self.client.post("/api/runs/synthetic-workflow-test/approve",
                                    headers={"Origin": "https://untrusted.example"}, json={"reviewer": "Test"})
        self.assertEqual(response.status_code, 403)

    def test_spreadsheet_formula_injection_is_escaped(self):
        self.run["run_id"] = "injection-test"
        self.run["rows"][0]["name"] = '=HYPERLINK("https://example.invalid","unsafe")'
        storage.save_run(self.run)
        response = self.client.get("/api/runs/injection-test/export?format=xlsx")
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        cell = workbook["Заказы поставщикам"]["E2"]
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("'="))
        workbook.close()

    def test_chart_details_are_loaded_separately(self):
        self.run["run_id"] = "detail-test"
        self.run["rows"][0]["history"] = [{"date": "2026-09-01", "quantity": 10}]
        storage.save_run(self.run)
        compact = self.client.get("/api/runs/detail-test").json()
        self.assertNotIn("history", compact["rows"][0])
        self.assertTrue(compact["rows"][0]["detail_available"])
        detail = self.client.get("/api/runs/detail-test/item", params={"sku": "TEST-001", "warehouse": "TEST-WH"})
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.json()["history"][0]["quantity"], 10)
        self.assertEqual(self.client.get("/api/runs/detail-test/item", params={"sku": "missing", "warehouse": "TEST-WH"}).status_code, 404)


if __name__ == "__main__":
    unittest.main()
