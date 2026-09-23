"""Black-box acceptance properties derived from the Elektrokоmplekt case.

These small synthetic fixtures test input-to-output behavior, not the engine's
internal formulas or the confidential partner archive.
"""

from copy import deepcopy
from datetime import date, timedelta
import csv
import io

from engine import calculate
from exports import COLUMNS, to_csv, to_xlsx
from openpyxl import load_workbook


TODAY = date(2026, 9, 23)


def dataset():
    start = TODAY - timedelta(days=119)
    return {
        "as_of": TODAY.isoformat(),
        "products": [
            {"sku": "SKU-1", "name": "Synthetic cable", "supplier": "Supplier A",
             "category": "Cable", "warehouse": "Almaty", "on_hand": 10,
             "lead_days": 7, "unit_price": 100, "pack_size": 1},
        ],
        "sales": [
            {"date": (start + timedelta(days=i)).isoformat(), "sku": "SKU-1",
             "warehouse": "Almaty", "quantity": 4, "client_id": "ANON-1"}
            for i in range(120)
        ],
        "stockouts": [], "inbound": [], "category_policies": {},
        "settings": {"review_days": 14, "safety_days": 7},
        "metadata": {"synthetic": True, "source_label": "Synthetic acceptance fixture"},
    }


def single_row(data):
    rows = calculate(data)["rows"]
    assert len(rows) == 1
    return rows[0]


def test_every_quantitative_must_have_input_changes_a_recommendation():
    baseline = dataset()
    base = single_row(baseline)["recommended_quantity"]

    higher_sales = deepcopy(baseline)
    for sale in higher_sales["sales"]:
        sale["quantity"] = 5
    assert single_row(higher_sales)["recommended_quantity"] > base

    higher_stock = deepcopy(baseline)
    higher_stock["products"][0]["on_hand"] += 40
    assert single_row(higher_stock)["recommended_quantity"] < base

    incoming = deepcopy(baseline)
    incoming["inbound"].append({"sku": "SKU-1", "warehouse": "Almaty",
                                "quantity": 40, "eta": (TODAY + timedelta(days=1)).isoformat()})
    assert single_row(incoming)["recommended_quantity"] < base

    sku_growth = deepcopy(baseline)
    sku_growth["products"][0]["growth"] = 0.25
    assert single_row(sku_growth)["recommended_quantity"] > base

    category_growth = deepcopy(baseline)
    category_growth["category_policies"]["Cable"] = {"growth": 0.25}
    assert single_row(category_growth)["recommended_quantity"] > base

    longer_lead = deepcopy(baseline)
    longer_lead["products"][0]["lead_days"] += 10
    assert single_row(longer_lead)["recommended_quantity"] > base


def test_same_sku_isolated_by_warehouse_and_grouped_by_supplier():
    data = dataset()
    other = deepcopy(data["products"][0])
    other.update(warehouse="Astana", supplier="Supplier B", on_hand=1)
    data["products"].append(other)
    second_sales = [dict(sale, warehouse="Astana", quantity=2)
                    for sale in data["sales"]]
    data["sales"].extend(second_sales)
    original = calculate(data)
    rows = {r["warehouse"]: r for r in original["rows"]}
    assert {r["supplier"] for r in original["rows"]} == {"Supplier A", "Supplier B"}
    assert {g["supplier"] for g in original["supplier_groups"]} == {"Supplier A", "Supplier B"}
    assert all(r["explanation"] and str(r["recommended_quantity"]) in r["explanation"]
               for r in original["rows"])

    data["inbound"] = [{"sku": "SKU-1", "warehouse": "Astana", "quantity": 40,
                        "eta": (TODAY + timedelta(days=1)).isoformat()}]
    revised = {r["warehouse"]: r for r in calculate(data)["rows"]}
    assert revised["Astana"]["recommended_quantity"] < rows["Astana"]["recommended_quantity"]
    assert revised["Almaty"]["recommended_quantity"] == rows["Almaty"]["recommended_quantity"]


def test_export_is_a_flat_reviewable_csv_and_xlsx_not_a_sent_order():
    result = calculate(dataset())
    result.update(run_id="acceptance-synthetic", source_label="Synthetic fixture", approval=None)
    csv_rows = list(csv.reader(io.StringIO(to_csv(result).decode("utf-8-sig")), delimiter=";"))
    assert csv_rows[0] == [title for _, title in COLUMNS]
    assert len(csv_rows) == 2
    assert "Черновик" in csv_rows[1][0]
    assert csv_rows[1][1:4] == ["Supplier A", "Almaty", "SKU-1"]
    assert csv_rows[1][11]

    workbook = load_workbook(io.BytesIO(to_xlsx(result)), read_only=True, data_only=True)
    try:
        sheet = workbook["Заказы поставщикам"]
        assert sheet.max_row == 2
        assert [cell.value for cell in sheet[1]] == [title for _, title in COLUMNS]
        assert workbook["Методология"]["B5"].value.startswith("Черновик")
    finally:
        workbook.close()
