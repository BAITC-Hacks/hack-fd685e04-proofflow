"""API-level acceptance checks tied to the five task requirements."""

from __future__ import annotations

import io
import os
from copy import deepcopy
from datetime import date, timedelta
from tempfile import TemporaryDirectory
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from server import app


AS_OF = date(2026, 9, 23)


def sample(*, days=90, quantity=4, on_hand=10):
    start = AS_OF - timedelta(days=days - 1)
    return {
        "as_of": AS_OF.isoformat(),
        "products": [{"sku": "SYNTH-A", "name": "Synthetic switch",
                      "supplier": "SYNTH-SUPPLIER", "category": "switchgear",
                      "warehouse": "SYNTH-W", "unit": "pcs", "on_hand": on_hand,
                      "lead_days": 7, "pack_size": 1, "unit_price": 10}],
        "sales": [{"date": (start + timedelta(days=i)).isoformat(),
                   "sku": "SYNTH-A", "warehouse": "SYNTH-W",
                   "quantity": quantity, "client_id": "SYNTH-REGULAR"}
                  for i in range(days)],
        "stockouts": [], "inbound": [],
        "category_policies": {},
        "settings": {"review_days": 14, "safety_days": 7},
        "metadata": {"synthetic": True, "source_label": "Synthetic acceptance fixture"},
    }


@pytest.fixture
def client():
    with TemporaryDirectory() as folder, patch.dict(os.environ, {"PROOFFLOW_DATA_DIR": folder}):
        with TestClient(app) as api:
            yield api


def calculate(client, data):
    response = client.post("/api/calculate", json={"dataset": data})
    assert response.status_code == 200, response.text[:500]
    return response.json()


def item(client, data):
    return calculate(client, data)["rows"][0]


def test_requirement_1_each_business_input_changes_recommendation(client):
    base = sample()
    original = item(client, base)
    stocked = deepcopy(base)
    stocked["products"][0]["on_hand"] += 70
    assert item(client, stocked)["recommended_quantity"] < original["recommended_quantity"]

    more_sales = deepcopy(base)
    for sale in more_sales["sales"]:
        sale["quantity"] = 8
    assert item(client, more_sales)["recommended_quantity"] > original["recommended_quantity"]

    arriving = deepcopy(base)
    arriving["inbound"] = [{"sku": "SYNTH-A", "warehouse": "SYNTH-W",
                            "quantity": 60, "eta": (AS_OF + timedelta(days=1)).isoformat()}]
    assert item(client, arriving)["recommended_quantity"] < original["recommended_quantity"]

    growing = deepcopy(base)
    growing["products"][0]["growth"] = 0.2
    assert item(client, growing)["recommended_quantity"] > original["recommended_quantity"]

    category = deepcopy(base)
    category["category_policies"]["switchgear"] = {"growth": 0.25, "safety_days": 12}
    assert item(client, category)["recommended_quantity"] > original["recommended_quantity"]


def test_requirement_2_calendar_seasonality_and_sustained_trend(client):
    seasonal = sample(days=730, quantity=1, on_hand=0)
    for sale in seasonal["sales"]:
        if date.fromisoformat(sale["date"]).month in (9, 10):
            sale["quantity"] = 5
    season = item(client, seasonal)
    assert season["seasonality_factor"] > 1

    trend = sample(days=90, quantity=1, on_hand=0)
    for index, sale in enumerate(trend["sales"]):
        sale["quantity"] = 1 + index // 18
    risen = item(client, trend)
    assert risen["trend_factor"] > 1
    assert risen["excluded_quantity"] == 0


def test_requirement_3_explicit_stockout_recovers_hidden_demand(client):
    data = sample(days=90, on_hand=0)
    start, end = AS_OF - timedelta(days=18), AS_OF - timedelta(days=7)
    for sale in data["sales"]:
        if start <= date.fromisoformat(sale["date"]) <= end:
            sale["quantity"] = 0
    without = item(client, data)
    with_interval = deepcopy(data)
    with_interval["stockouts"] = [{"sku": "SYNTH-A", "warehouse": "SYNTH-W",
                                   "start": start.isoformat(), "end": end.isoformat()}]
    adjusted = item(client, with_interval)
    assert adjusted["lost_demand"] > 0
    assert adjusted["recommended_quantity"] > without["recommended_quantity"]


def test_requirement_4_one_off_order_does_not_inflate_normal_demand(client):
    data = sample(days=90)
    base = item(client, data)
    spike = deepcopy(data)
    spike["sales"].append({"date": (AS_OF - timedelta(days=4)).isoformat(),
                           "sku": "SYNTH-A", "warehouse": "SYNTH-W",
                           "quantity": 500, "client_id": "SYNTH-ONE-OFF"})
    cleaned = item(client, spike)
    assert cleaned["excluded_quantity"] >= 500
    assert cleaned["recommended_quantity"] == base["recommended_quantity"]


def test_requirement_5_supplier_groups_explanation_approval_and_export(client):
    demo = client.post("/api/demo")
    assert demo.status_code == 200
    assert demo.json()["dataset"]["metadata"]["synthetic"] is True
    result = client.post("/api/calculate", json={})
    assert result.status_code == 200, result.text[:500]
    body = result.json()
    assert body["supplier_groups"]
    assert sum(g["order_lines"] for g in body["supplier_groups"]) == body["summary"]["order_lines"]
    chosen = next(row for row in body["rows"] if row["recommended_quantity"] > 0)
    assert chosen["explanation"]
    assert chosen["supplier"]
    detail = client.get(f"/api/runs/{body['run_id']}/item",
                        params={"sku": chosen["sku"], "warehouse": chosen["warehouse"]})
    assert detail.status_code == 200
    assert detail.json()["breakdown"]["order_rounding"]["recommended"] == chosen["recommended_quantity"]

    key = chosen["warehouse"] + "::" + chosen["sku"]
    edited = float(chosen["recommended_quantity"] + 5)
    approved = client.post(f"/api/runs/{body['run_id']}/approve",
                           json={"reviewer": "Synthetic acceptance reviewer",
                                 "quantities": {key: edited}})
    assert approved.status_code == 200, approved.text[:500]
    assert approved.json()["quantities"][key] == edited
    exported = client.get(f"/api/runs/{body['run_id']}/export?format=xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content), data_only=True)
    try:
        assert workbook["Заказы поставщикам"].max_row > 1
    finally:
        workbook.close()
